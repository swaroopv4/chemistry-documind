"""Real generated Office files and chemistry fixtures; no paid service calls."""
from pathlib import Path
import struct
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch
from io import BytesIO

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core.formats import extract_sections, Section
from core.ingestion import ingest_file, chunk_sections
from core import legacy_ppt

MOL2 = """@<TRIPOS>MOLECULE
water
3 2 0 0 0
SMALL
USER_CHARGES
@<TRIPOS>ATOM
1 O1 0.0 0.0 0.0 O.3 1 WAT -0.8
2 H1 0.8 0.0 0.0 H 1 WAT 0.4
3 H2 -0.2 0.8 0.0 H 1 WAT 0.4
@<TRIPOS>BOND
1 1 2 1
2 1 3 1
"""

CIF = """data_water
_chemical_formula_sum 'H2 O'
_cell_length_a 9.12(3)
_description
;
line one
line two
;
loop_
_atom_site_label
_atom_site_fract_x
_atom_site_fract_y
O1 0.0 ?
H1 0.1 .
save_details
_details.note 'quoted note'
save_
data_second
_chemical_formula_sum 'C H4'
"""


class FormatChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_pptx_text_tables_and_speaker_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = 'Chemistry project'
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
        table.cell(0, 0).text = 'Atom'
        table.cell(0, 1).text = 'Charge'
        table.cell(1, 0).text = 'O1'
        table.cell(1, 1).text = '-0.8'
        slide.notes_slide.notes_text_frame.text = 'Charges are partial charges'
        path = self.root / 'slides.pptx'
        deck.save(str(path))
        sections = extract_sections(path)
        self.assertIn('O1 | -0.8', sections[0].text)
        self.assertEqual(sections[0].locator, 'slide 1')
        self.assertIn('speaker notes', sections[1].locator)
        chunks = ingest_file(path)
        self.assertTrue(all(c.doc_name == path.name and c.metadata['file_type'] == 'pptx' for c in chunks))

    def test_docx_preserves_body_order_tables_and_headers(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Before table')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Parameter'
        table.cell(0, 1).text = 'Value'
        table.cell(1, 0).text = 'pH'
        table.cell(1, 1).text = '7.4'
        document.add_paragraph('After table')
        document.sections[0].header.paragraphs[0].text = 'Lab report'
        path = self.root / 'report.docx'
        document.save(str(path))
        sections = extract_sections(path)
        self.assertEqual([s.text for s in sections[:4]],
                         ['Before table', 'Parameter | Value', 'pH | 7.4', 'After table'])
        self.assertIn('table row 2', sections[2].locator)
        self.assertTrue(any(s.text == 'Lab report' for s in sections))

    def test_mol2_multiple_molecules_row_context_and_validation(self):
        path = self.root / 'molecules.mol2'
        path.write_text(MOL2 + MOL2.replace('water', 'water2'), encoding='utf-8')
        chunks = ingest_file(path)
        self.assertTrue(any(c.metadata['locator'].startswith('molecule 2') for c in chunks))
        atom_chunks = [c for c in chunks if '@<TRIPOS>ATOM' in c.text]
        self.assertEqual(len(atom_chunks), 2)
        self.assertIn('partial_charge', atom_chunks[0].text)
        self.assertIn('0.0 0.0 0.0 O.3', atom_chunks[0].text)
        for malformed in [MOL2.replace('3 2 0', '4 2 0'), MOL2.replace('2 1 3 1', '2 1 9 1'),
                          MOL2.replace('0.8 0.0 0.0', 'nan 0.0 0.0')]:
            path.write_text(malformed, encoding='utf-8')
            with self.assertRaises(ValueError):
                ingest_file(path)

    def test_cif_quotes_uncertainties_missing_values_loops_and_blocks(self):
        path = self.root / 'crystal.cif'
        path.write_text(CIF, encoding='utf-8')
        sections = extract_sections(path)
        text = '\n'.join(s.prefix + '\n' + s.text for s in sections)
        self.assertIn('9.12(3)', text)
        self.assertIn('line one\\nline two', text)
        self.assertIn('_atom_site_fract_y=?', text)
        self.assertIn('_atom_site_fract_y=.', text)
        self.assertTrue(any('save_details' in s.locator for s in sections))
        self.assertTrue(any(s.locator.startswith('data_second') for s in sections))
        chunks = ingest_file(path)
        self.assertTrue(all(c.token_count <= 400 for c in chunks))
        path.write_text('data_bad\nloop_\n_tag1\n_tag2\nonly_one_value', encoding='utf-8')
        with self.assertRaises(ValueError):
            ingest_file(path)

    def test_chemical_chunking_keeps_rows_and_repeats_identifiers(self):
        rows = [f'row {i}: _atom_site_label=O{i} | _atom_site_fract_x=0.25' for i in range(50)]
        section = Section('\n'.join(rows), 'data_test, loop _atom_site_label', 1, 'data_test\n_cell_length_a=9.1')
        chunks = chunk_sections([section], 'test.cif', 'cif', 100, 10)
        self.assertGreater(len(chunks), 1)
        collected = [row for c in chunks for row in c.text.splitlines()[2:]]
        self.assertEqual(collected, rows)
        self.assertTrue(all(c.text.startswith(section.prefix) and c.token_count <= 100 for c in chunks))

    def test_unsupported_blank_and_invalid_chunk_settings(self):
        path = self.root / 'file.txt'
        path.write_text('text')
        with self.assertRaisesRegex(ValueError, 'Unsupported file type'):
            ingest_file(path)
        from docx import Document
        path = self.root / 'blank.docx'
        Document().save(path)
        with self.assertRaisesRegex(ValueError, 'No searchable text'):
            ingest_file(path)
        with self.assertRaises(ValueError):
            chunk_sections([Section('text', 'item 1', 1)], 'x.docx', 'docx', 100, 100)

    def test_legacy_ppt_uses_latest_edits_and_skips_deleted_slides(self):
        # Construct a real MS-PPT record stream. OLE transport is mocked here;
        # extraction is separately verified with Apache POI's binary PPT fixture.
        def rec(kind, payload, flags=0):
            return struct.pack('<HHI', flags, kind, len(payload)) + payload
        def slide(text):
            textbox = rec(0xF00D, rec(4000, text.encode('utf-16-le')))
            return rec(1006, textbox, 15)
        def doc():
            outline = rec(1011, struct.pack('<IIIII', 2, 0, 1, 256, 0))
            return rec(1000, rec(4080, outline, 15), 15)
        data = bytearray()
        old_doc = len(data); data += doc()
        old_slide = len(data); data += slide('Obsolete saved content')
        deleted = len(data); data += slide('Deleted slide content')
        old_dir = len(data)
        data += rec(6002, struct.pack('<IIII', (3 << 20) | 1, old_doc, old_slide, deleted))
        old_edit = len(data)
        data += rec(4085, struct.pack('<IIIIIII', 256, 0, 0, old_dir, 1, 4, 0))
        new_slide = len(data); data += slide('Latest chemistry results')
        new_dir = len(data); data += rec(6002, struct.pack('<II', (1 << 20) | 2, new_slide))
        new_edit = len(data)
        data += rec(4085, struct.pack('<IIIIIII', 256, 0, old_edit, new_dir, 1, 4, 0))
        user = rec(4086, struct.pack('<III', 20, 0xE391C05F, new_edit))
        ole = MagicMock()
        ole.__enter__.return_value = ole
        ole.openstream.side_effect = lambda name: BytesIO(bytes(data) if name == 'PowerPoint Document' else user)
        with patch.object(legacy_ppt.olefile, 'isOleFile', return_value=True), \
             patch.object(legacy_ppt.olefile, 'OleFileIO', return_value=ole):
            sections = legacy_ppt.extract_ppt(Path('test.ppt'))
        self.assertEqual([(s.locator, s.text) for s in sections], [('slide 1', 'Latest chemistry results')])


if __name__ == '__main__':
    unittest.main(verbosity=2)
