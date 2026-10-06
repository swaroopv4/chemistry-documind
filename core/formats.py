"""Extract searchable text with honest locations and optional local PDF OCR."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class Section:
    text: str
    locator: str
    number: int
    # A short identifying prefix is repeated when a chemical section spans chunks.
    prefix: str = ""
    ocr: bool = False
    ocr_confidence: float = 0.


SUPPORTED_EXTENSIONS = ("pdf", "ppt", "pptx", "docx", "mol2", "cif")


def extract_sections(path: Path) -> list[Section]:
    extension = path.suffix.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    if extension == "pdf":
        from core.ingestion import extract_text_from_pdf
        ocr_results = {}
        return [Section(text, f"page {number}", number,ocr=number in ocr_results,ocr_confidence=ocr_results.get(number,0.))
                for number, text in extract_text_from_pdf(path,ocr_results)]
    if extension == "pptx":
        return extract_pptx(path)
    if extension == "ppt":
        from core.legacy_ppt import extract_ppt
        return extract_ppt(path)
    if extension == "docx":
        return extract_docx(path)
    from core.chemical_formats import extract_cif, extract_mol2
    return extract_mol2(path) if extension == "mol2" else extract_cif(path)


def extract_pptx(path: Path) -> list[Section]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shape_text(shape.shapes)
            if shape.has_text_frame and shape.text.strip():
                yield shape.text
            if shape.has_table:
                for row in shape.table.rows:
                    yield " | ".join(cell.text for cell in row.cells)

    result = []
    for number, slide in enumerate(Presentation(str(path)).slides, 1):
        text = "\n".join(shape_text(slide.shapes))
        if text.strip():
            result.append(Section(text, f"slide {number}", number))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes and notes.text.strip():
                result.append(Section(notes.text, f"slide {number}, speaker notes", number))
    return result


def extract_docx(path: Path) -> list[Section]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    result = []
    document = Document(str(path))
    # XML body order preserves the order of paragraphs and tables. Pagination is
    # layout-dependent, so Word citations use body item numbers rather than pages.
    def visit(container, parent, location):
        for number, element in enumerate(container, 1):
            locator = f"{location} {number}"
            if element.tag.endswith("}p"):
                text = Paragraph(element, parent).text
                if text.strip():
                    result.append(Section(text, locator + ", paragraph", len(result) + 1))
            elif element.tag.endswith("}tbl"):
                table = Table(element, parent)
                for row_num, row in enumerate(table.rows, 1):
                    cells = []
                    seen = set()
                    for cell in row.cells:
                        if cell._tc in seen:
                            continue  # merged cells
                        seen.add(cell._tc)
                        cells.append(cell.text)
                        for nested in cell.tables:
                            visit([nested._tbl], cell, f"{locator}, row {row_num}, nested item")
                    text = " | ".join(cells)
                    if text.strip(" |"):
                        result.append(Section(text, f"{locator}, table row {row_num}", len(result) + 1))

    visit(document.element.body, document, "body item")
    for section_num, section in enumerate(document.sections, 1):
        for kind in ("header", "first_page_header", "even_page_header",
                     "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, kind)
            if not part.is_linked_to_previous:
                visit(part._element, part, f"section {section_num}, {kind} item")
    return result
